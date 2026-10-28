"""Every slide's words in one Word document, for rewriting the text elsewhere.

    uv run --with beautifulsoup4 --with python-docx python tool/slides/export_text.py OUT_DIR
    # OUT_DIR/talk_text.docx

Reads the built slides.html and script.html. Per slide: title, message, what else the slide
shows, its figure labels, speaker notes and, for timed slides, the spoken script and the bridge
to the next slide. Numbers appear as rendered; they come from the data and must not change.
"""

from __future__ import annotations

import argparse
import re
from pathlib import Path

from bs4 import BeautifulSoup, Tag
from docx import Document
from docx.shared import Pt

HERE = Path(__file__).resolve().parent
BLOCKS = (
    "div.stat, div.cond, div.case-row, div.agree-row, div.cmp-row, div.rule, tr, li, "
    "blockquote, dt, figcaption, h3, p"
)
#: Blocks whose parts read as one line, joined by a separator.
JOINED = {"stat", "cond", "case-row", "agree-row", "cmp-row", "rule"}
BRIEF = (
    "This document holds the words of the talk “Reading What Isn’t There”, slide by slide, so "
    "they can be rewritten. Keep every number, unit, name and citation exactly as written: the "
    "numbers are computed from the data and must not change or be rounded. Titles state a "
    "finding in plain English, one idea each, with no meta wording (“this talk”, “my abstract”, "
    "“the goal”). Act III is microplastics, counted in public datasets; Act IV is nanoplastics, "
    "lab particles measured for charge in papers; keep the two apart. A spoken script must fit "
    "its slide’s time at no more than 130 words a minute. Slides 1 to 23 are timed; the rest "
    "are backup for questions. Return the same structure, field by field."
)


def words(el: Tag) -> str:
    return re.sub(r"\s+", " ", el.get_text()).strip()


def block_text(el: Tag) -> str:
    cls = set(el.get("class") or [])
    if cls & JOINED:
        parts = [words(c) for c in el.find_all(recursive=False) if isinstance(c, Tag)]
        return " · ".join(p for p in parts if p)
    if el.name == "tr":
        return " | ".join(words(c) for c in el.find_all(["td", "th"]))
    if el.name == "dt":
        dd = el.find_next_sibling("dd")
        return f"{words(el)}: {words(dd)}" if dd else words(el)
    if el.name == "blockquote":
        quote, cite = el.find("p"), el.find("cite")
        return words(quote) + (f" ({words(cite)})" if cite else "")
    return words(el)


def on_slide(body: Tag | None) -> list[str]:
    """The slide's own text in reading order, each block once, nothing from inside a figure."""
    if body is None:
        return []
    for svg in body.find_all("svg"):
        svg.extract()
    chosen = body.select(BLOCKS)
    keep = [e for e in chosen if not any(p in chosen for p in e.parents)]
    return [t for t in (block_text(e) for e in keep) if t]


def figure_labels(section: Tag) -> list[str]:
    seen: list[str] = []
    for t in section.select("svg text"):
        s = words(t)
        if re.search(r"[A-Za-z]{2}", s) and s not in seen:
            seen.append(s)
    return seen


def slides(deck: Path) -> list[dict]:
    soup = BeautifulSoup(deck.read_text(encoding="utf-8"), "html.parser")
    out = []
    for sec in soup.select("section.slide"):
        msg, notes = sec.select_one(".msg"), sec.select_one(".notes p")
        eyebrow = sec.select_one(".eyebrow span:last-child")
        labels = figure_labels(sec)
        header = [words(p) for p in sec.select("header p")
                  if not {"eyebrow", "msg"} & set(p.get("class") or [])]  # fmt: skip
        out.append({
            "id": sec["id"],
            "backup": "backup" in sec["class"],
            "act": words(eyebrow) if eyebrow else "",
            "title": words(sec.select_one("h1, h2")),
            "message": words(msg) if msg else "",
            "shown": header + on_slide(sec.select_one(".body")),
            "figure": labels,
            "notes": words(notes) if notes else "",
        })  # fmt: skip
    return out


def script(path: Path) -> dict[str, dict]:
    soup = BeautifulSoup(path.read_text(encoding="utf-8"), "html.parser")
    out = {}
    for card in soup.select("section.card"):
        for pause in card.select(".pause"):
            pause.replace_with("(pause)")
        time = card.select_one(".time b")
        then = card.select_one(".cue .then")
        if then and then.select_one(".k"):
            then.select_one(".k").extract()
        out[card["id"]] = {
            "time": words(time) if time else "",
            "say": [words(p) for p in card.select(".say p")],
            "then": words(then) if then else "",
        }
    return out


def field(doc: Document, label: str, text: str) -> None:
    p = doc.add_paragraph()
    p.add_run(label + ": ").bold = True
    p.add_run(text)


def write(deck: list[dict], spoken: dict[str, dict], path: Path) -> None:
    doc = Document()
    doc.styles["Normal"].font.size = Pt(11)
    doc.add_heading("Reading What Isn’t There: the text of every slide", 0)
    doc.add_paragraph(BRIEF)
    for s in deck:
        n = s["id"][1:]
        sp = spoken.get(s["id"], {})
        kind = "backup" if s["backup"] else f"timed, {sp.get('time', '')}"
        doc.add_heading(f"Slide {n} · {s['act']} ({kind})", 1)
        field(doc, "Title", s["title"])
        if s["message"]:
            field(doc, "Message", s["message"])
        if s["shown"]:
            doc.add_paragraph().add_run("On the slide:").bold = True
            for t in s["shown"]:
                doc.add_paragraph(t, style="List Bullet")
        if s["figure"]:
            field(doc, "Figure labels", " · ".join(s["figure"]))
        if s["notes"]:
            field(doc, "Speaker notes", s["notes"])
        if not s["backup"] and sp.get("say"):
            doc.add_paragraph().add_run("Spoken script:").bold = True
            for t in sp["say"]:
                doc.add_paragraph(t)
            if sp.get("then"):
                field(doc, "Bridge to the next slide", sp["then"])
    doc.save(path)


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("out", type=Path)
    ap.add_argument("--deck", type=Path, default=HERE / "slides.html")
    ap.add_argument("--script", type=Path, default=HERE / "script.html")
    a = ap.parse_args()
    a.out.mkdir(parents=True, exist_ok=True)
    deck = slides(a.deck)
    write(deck, script(a.script), a.out / "talk_text.docx")
    print(f"{len(deck)} slides -> {a.out / 'talk_text.docx'}")


if __name__ == "__main__":
    main()
