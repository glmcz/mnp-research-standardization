"""Second routes to P2's numbers: each is recomputed here without the builder's pandas code.

uv run pytest -q -p no:cacheprovider tool/p2/test_p2.py   (run tool/p2/build.py first)
"""

from __future__ import annotations

import csv
import json
import re
import statistics
import sys
from pathlib import Path

import pytest

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))
import build  # noqa: E402

DATA = json.loads((HERE / "data.json").read_text(encoding="utf-8"))


def ranks(xs: list[float]) -> list[float]:
    order = sorted(range(len(xs)), key=xs.__getitem__)
    out = [0.0] * len(xs)
    i = 0
    while i < len(order):
        j = i
        while j + 1 < len(order) and xs[order[j + 1]] == xs[order[i]]:
            j += 1
        for k in range(i, j + 1):
            out[order[k]] = (i + j) / 2 + 1
        i = j + 1
    return out


def rho(x: list[float], y: list[float]) -> float:
    return statistics.correlation(ranks(x), ranks(y))


def test_inputs_match_their_pins():
    build.check_pins()


def test_amine_polystyrene_signs_from_the_d1_table():
    neg = pos = 0
    with open(build.D1, encoding="utf-8") as f:
        for r in csv.DictReader(f):
            size = r["size_nominal_nm"] or r["size_measured_nm"]
            if not (
                r["quantity"] == "zeta_potential"
                and r["unit"] == "mV"
                and r["value"]
                and size
                and float(size) > 0
            ):
                continue
            if "polystyren" not in r["material"].lower() or not re.search(
                r"amin|nh2|nh3", r["functionalisation"].lower()
            ):
                continue
            v = float(r["value"])
            neg += v < 0
            pos += v > 0
    s = DATA["zeta"]["stats"]
    assert (neg, pos) == (s["amine_negative"], s["amine_positive"])


def test_ncei_records_per_setting_from_the_raw_csv():
    counts: dict[str, int] = {}
    with open(build.NCEI, encoding="utf-8-sig") as f:
        for r in csv.DictReader(f):
            counts[r["Marine Setting"]] = counts.get(r["Marine Setting"], 0) + 1
    for setting, v in DATA["ncei"]["stats"]["by_setting"].items():
        assert counts[setting] == v["records"], setting


def test_finer_limit_higher_count_from_the_raw_csv():
    per: dict[str, tuple[list[float], list[float]]] = {}
    with open(build.NCEI, encoding="utf-8-sig") as f:
        for r in csv.DictReader(f):
            if r["Marine Setting"] != "Ocean water":
                continue
            mesh, vals = per.setdefault(r["DOI"].strip(), ([], []))
            if r["Mesh Size (mm)"]:
                mesh.append(float(r["Mesh Size (mm)"]))
            try:
                v = float(r["Microplastics Measurement"])
            except ValueError:
                continue
            if v > 0:
                vals.append(v)
    rows = [(min(m), statistics.median(v)) for m, v in per.values() if m and v]
    want = DATA["ncei"]["stats"]["water_lower_vs_median"]
    assert len(rows) == want["n"]
    assert rho([a for a, _ in rows], [b for _, b in rows]) == pytest.approx(want["rho"], abs=0.005)


def test_case_ratios_from_the_raw_csv():
    """The tow pairing comes from the case study; the values are read again from NCEI itself."""
    value = {}
    with open(build.NCEI, encoding="utf-8-sig") as f:
        for r in csv.DictReader(f):
            try:
                value[r["Unique ID"]] = float(r["Microplastics Measurement"])
            except ValueError:
                continue
    by_method: dict[str, list[float]] = {}
    with open(build.CASE_TOWS, encoding="utf-8") as f:
        tows = list(csv.DictReader(f))
    for t in tows:
        a, b = value[t["Unique ID_R"]], value[t["Unique ID_E"]]
        if a > 0 and b > 0:
            by_method.setdefault(t["Sampling Method"], []).append(a / b)
    c = DATA["case"]
    assert len(tows) == c["tows"]
    for method, ratios in by_method.items():
        assert round(statistics.median(ratios), 2) == c["median_ratio_by_method"][method]


def test_land_outline_never_jumps_across_the_map():
    """Rings crossing the date line used to draw a line across the whole map."""
    for sub in DATA["land"].split("M")[1:]:
        xs = [float(p.split(",")[0]) for p in sub.rstrip("Z").split("L")]
        assert (
            max(abs(b - a) for a, b in zip(xs, xs[1:], strict=False)) < 180 if len(xs) > 1 else True
        )


def test_page_prose_types_no_numbers():
    """Numbers on the page come from DATA. Typed digits allowed: dates and names of sources."""
    html = (HERE / "template.html").read_text(encoding="utf-8")
    prose = re.sub(r"<(script|style)\b.*?</\1>", "", html, flags=re.S)
    prose = re.sub(r"<[^>]+>", " ", prose)
    prose = re.sub(r"\bPart \d\b|\bP2\b", " ", prose)  # section labels, not data
    allowed = {"2026-10-10", "10th", "90th"}
    typed = [t for t in re.findall(r"\S*\d\S*", prose) if t.strip(".,;:()") not in allowed]
    assert typed == []


def test_polymer_medians_from_the_saved_table():
    by: dict[str, list[float]] = {}
    with open(build.ROOT / "out/tables/p2_zeta_size.csv", encoding="utf-8") as f:
        for r in csv.DictReader(f):
            by.setdefault(r["material_group"], []).append(float(r["zeta_mV"]))
    for p in DATA["zeta"]["stats"]["by_polymer"]:
        assert len(by[p["polymer"]]) == p["rows"]
        assert round(statistics.median(by[p["polymer"]]), 1) == p["zeta_median"], p["polymer"]


# ---- extension layer and pre-registered results (run with: uv run --with openpyxl pytest ...)
import extension as ext  # noqa: E402

PREREG = json.loads((build.ROOT / "out/tables/p2_prereg_results.json").read_text(encoding="utf-8"))


def test_extension_inputs_match_their_pins():
    ext.check_pins()


def test_page_shows_the_results_file_the_runner_wrote():
    assert DATA["ext"]["prereg"] == PREREG


def test_weser_windows_and_ratio_from_the_raw_files():
    def read(path):
        text = path.read_text(encoding="utf-8")
        head, body = text.split("*/\n", 1)
        window = re.search(r"Citation:\t.*?\((?:S|L)-MP, (\d+)-(\d+) µm\)", head).groups()
        rows = list(csv.DictReader(body.splitlines(), delimiter="\t"))
        return [float(w) for w in window], {
            r["Event"].split("_")[-1]: float(r["Microplastic [#/m**3]"]) for r in rows
        }

    ws, small = read(ext.WESER_S)
    wl, large = read(ext.WESER_L)
    ratios = [small[k] / large[k] for k in small if k in large and large[k] > 0]
    w = DATA["ext"]["weser"]
    assert (ws, wl) == (w["windows"]["small_um"], w["windows"]["large_um"])
    assert len(ratios) == len(w["pairs"]) - w["stations_without_large"]
    assert round(statistics.median(ratios), 1) == w["median_ratio"]


def test_implied_alpha_inverts_the_window_ratio():
    w = DATA["ext"]["weser"]["windows"]
    for r in (12.0, 800.0, 5340.4):
        a = ext.implied_alpha(r, w["small_um"], w["large_um"])
        assert ext.window_ratio(a, w["small_um"], w["large_um"]) == pytest.approx(r, rel=1e-6)


def test_p1_rho_from_the_soil_workbook():
    from openpyxl import load_workbook

    ws = load_workbook(ext.SOIL, read_only=True)["information"]
    rows = ws.iter_rows(values_only=True)
    head = list(next(rows))
    col = {name: head.index(name) for name in ("ID", "Title", "Size min(μm)", "abundance_numeric")}
    lower: dict[str, float] = {}
    conc: dict[str, list[float]] = {}
    for r in rows:
        if r[col["ID"]] is None:
            continue
        study = str(r[col["Title"]]).strip()
        lo, c = r[col["Size min(μm)"]], r[col["abundance_numeric"]]
        if isinstance(lo, (int, float)):
            lower[study] = min(lower.get(study, float("inf")), float(lo))
        if isinstance(c, (int, float)) and c > 0:
            conc.setdefault(study, []).append(float(c))
    pts = [(lower[s], statistics.median(conc[s])) for s in lower if s in conc and lower[s] > 0]
    import math

    assert len(pts) == PREREG["P1"]["studies"]
    got = rho([math.log10(a) for a, _ in pts], [math.log10(b) for _, b in pts])
    assert got == pytest.approx(PREREG["P1"]["rho"], abs=0.005)


def test_each_verdict_follows_its_committed_rule():
    p1, p2, p3, p4 = (PREREG[k] for k in ("P1", "P2", "P3", "P4"))
    lo, hi = p1["slope_ci95"]
    if p1["rho"] >= 0 or hi < -1.2 or lo > -0.3:
        want1 = "refuted"
    elif -0.40 <= p1["rho"] <= -0.15 and -1.2 <= p1["slope"] <= -0.3:
        want1 = "supported"
    else:
        want1 = "partly"
    m = p2["median_reduction"]
    want2 = "supported" if m < 0.25 else "refuted" if m >= 0.50 else "partly"
    within = p3["share_within_1.6pm0.3"]
    ok3 = (
        1.3 <= p3["median_alpha"] <= 2.2
        and 1 - within >= 0.30
        and p3["mannwhitney_p_two_largest_types"] < 0.05
    )
    want3 = "refuted" if within >= 0.80 else "supported" if ok3 else "partly"
    r = {k: v["r2"] for k, v in p4["feature_sets"].items()}
    refuted4 = r["all"] >= 0.69 or max(r["polymer"], r["size"]) > max(r["surface"], r["liquid_pH"])
    ok4 = (
        r["all"] < 0.5
        and r["size"] <= 0.05
        and r["surface"] > r["polymer"]
        and r["liquid_pH"] > r["polymer"]
    )
    want4 = "refuted" if refuted4 else "supported" if ok4 else "partly"
    assert [p1["verdict"], p2["verdict"], p3["verdict"], p4["verdict"]] == [
        want1,
        want2,
        want3,
        want4,
    ]
