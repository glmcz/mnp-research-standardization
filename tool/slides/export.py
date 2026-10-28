"""The deck as PDF and PPTX: each slide rendered from slides.html, its speaker notes kept.

    uv run --with playwright --with pillow --with python-pptx python tool/slides/export.py OUT_DIR
    # OUT_DIR/talk_deck.pdf, OUT_DIR/talk_deck.pptx

Each slide is a picture of the page as presented (16:9, rendered at twice 1280 x 720), so the
layout is exact and the slide text is not editable. The PPTX carries every slide's speaker notes
in PowerPoint's notes pane. Rebuild the deck first; this reads the built slides.html.
"""

from __future__ import annotations

import argparse
import io
import tempfile
from pathlib import Path

from PIL import Image
from playwright.sync_api import sync_playwright
from pptx import Presentation
from pptx.util import Inches

HERE = Path(__file__).resolve().parent
CHROME = "/Applications/Google Chrome.app/Contents/MacOS/Google Chrome"


def render(deck: Path) -> list[tuple[str, bytes, str, str]]:
    """(slide id, PNG, title, speaker notes) for every slide, in deck order."""
    with tempfile.TemporaryDirectory() as tmp:
        wrap = Path(tmp) / "deck.html"
        wrap.write_text(
            "<!doctype html><html><head><meta charset='utf-8'><meta name='viewport' "
            "content='width=device-width,initial-scale=1'></head><body>"
            + deck.read_text(encoding="utf-8")
            + "</body></html>",
            encoding="utf-8",
        )
        with sync_playwright() as p:
            b = p.chromium.launch(executable_path=CHROME)
            pg = b.new_page(
                viewport={"width": 1280, "height": 720}, device_scale_factor=2, color_scheme="light"
            )
            pg.goto(wrap.as_uri(), wait_until="networkidle")
            pg.add_style_tag(content=".controls { display: none !important; }")
            ids = pg.eval_on_selector_all("section.slide", "els => els.map(e => e.id)")
            out = []
            for sid in ids:
                pg.goto(wrap.as_uri() + "#" + sid)
                pg.add_style_tag(content=".controls { display: none !important; }")
                pg.wait_for_timeout(450)
                title, notes = pg.eval_on_selector(
                    f"#{sid}",
                    """s => [s.querySelector('h1, h2'), s.querySelector('.notes p')].map(
                        e => e ? e.textContent.replace(/\\s+/g, ' ').trim() : '')""",
                )
                out.append((sid, pg.screenshot(), title, notes))
            b.close()
    return out


def write_pdf(slides: list[tuple[str, bytes, str, str]], path: Path) -> None:
    pages = [Image.open(io.BytesIO(png)).convert("RGB") for _, png, _, _ in slides]
    pages[0].save(path, save_all=True, append_images=pages[1:], resolution=192.0, quality=92)


def write_pptx(slides: list[tuple[str, bytes, str, str]], path: Path) -> None:
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    blank = prs.slide_layouts[6]
    for sid, png, title, notes in slides:
        s = prs.slides.add_slide(blank)
        pic = s.shapes.add_picture(io.BytesIO(png), 0, 0, prs.slide_width, prs.slide_height)
        pic.name = f"slide {sid[1:]}"
        pic._element.nvPicPr.cNvPr.set("descr", title)
        if notes:
            s.notes_slide.notes_text_frame.text = notes
    prs.save(path)


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("out", type=Path)
    ap.add_argument("--deck", type=Path, default=HERE / "slides.html")
    a = ap.parse_args()
    a.out.mkdir(parents=True, exist_ok=True)
    slides = render(a.deck)
    write_pdf(slides, a.out / "talk_deck.pdf")
    write_pptx(slides, a.out / "talk_deck.pptx")
    print(f"{len(slides)} slides, {sum(bool(n) for *_, n in slides)} with notes -> {a.out}")


if __name__ == "__main__":
    main()
