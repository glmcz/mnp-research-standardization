"""Koelmans et al. 2019 (Water Res 155:410, PMC6449537) SI Table S1 parser.

The fixture docx mirrors the real `mmc1.docx`: the same 12 header strings, free-text
cells, and micrometre signs written as Symbol-font `w:sym` glyphs (char F06D), which
python-docx's `cell.text` silently drops ("333 µm mesh" reads as "333 m mesh").
"""

import io
import zipfile

import httpx
import numpy as np
import pandas as pd
import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from mnp.rescale import koelmans2019 as k

REAL_HEADERS = [
    "Reference",
    "Country",
    "Source",
    "Treatment",
    "Sampling Date",
    "Size, shape",
    "Polymers, chemicals",
    "Value",
    "Detection limit, negative and positive controls and blanks.  ",
    "Sampling method",
    "Analysis method",
    "Comments",
]

# Cells copied verbatim from the real Table S1; "{mu}" marks a Symbol-font µ glyph.
FIXTURE_ROWS = [
    [
        "Anderson et al. 2017",
        "Lake Winnipeg, Canada",
        "LAK",
        "n.a. ",
        "Jul 2014 - Jun 2016",
        "Fibres, fragments, film and foam (pellets not found).",
        "n.i.",
        "Mean: 193 420 (±115 567 SD) #/km2\nRange: 52 508 – 748 027 #/km2",
        "DI water blanks (quadruplicates, 480L). Air blanks for 24h. Corrected for both blanks. ",
        "Manta trawl; 333 {mu}m mesh. Preserved in 70% ethanol.  ",
        "Samples rinsed and large objects removed; WPO treatment with Fe(II), heated to 75 °C; "
        "visual inspection of subsamples for plastics; subset of particles identified with "
        "SEM-EDS. ",
        "\xa0",
    ],
    [
        "Baldwin et al., 2016\n",
        "Great Lake tributaries, US\n",
        "RIV",
        "n.a.",
        "Apr 2014 - Apr 2015\n",
        "Size: 0.355-0.999, 1.00-4.759, ≥4.75 mm; \n\nShape: Fragments, pellets/beads, "
        "lines/fibers, films and foam.\n",
        "n.i.",
        "Mean: 4.2 x 10-3 #/L;\nMedian: 1.9 x 10-3 #/”L;\nRange: 0.05 - 32 (x10-3) #/L \n",
        "Five negative controls in the field and 11 in the lab were included. ",
        "Neuston net (333 µm). Sample volume measured. . Net rinsed with tap water or filtered "
        "(333 µm) stream water. Preserved in isopropyl alcohol.",
        "Sieving through 4.75, 1.00 and 0.355 mm mesh. WPO with Fe(II) catalyst at 75°C). "
        "WPO solution sieved through 125 µm and MP visually identified under dissection "
        "microscope (40x). \n",
        "",
    ],
    [
        "Mason et al. 2018",
        "US and others",
        "BOT",
        "n.a.",
        "2017",
        "Size: 6.5-100, >100 µm;\nShape: Fragment, film, fibre, foam, pellet",
        "PP, nylon, PS, PE, PEST",
        "Mean: 325 #/L (> 100 {mu}m - 10.4 #/L, < 100{mu}m - 315 #/L);  \nRange: 0-10390 #/L",
        "Lab blanks included.",
        "Bottles opened in laminar flow hood and Nile Red dye added.",
        "Filtered on glass fibre filter; fluorescence microscopy; FTIR on a subset.",
        "",
    ],
    [
        "Faure et al. 2015",
        "Lake Geneva and rivers, CH",
        "LAK, RIV",
        "n.a.",
        "2013",
        "Size: 300 – 5000 µm.  \n\nShape:Fragments, pellets, beads, lines, fibres, films, foams.",
        "PE, PP, PS",
        "Lakes: Mean 91 000 (±120 000) #/km-2 or 26 000 (33 000) mg/km2.  \n\nMedian: 48 000 "
        "#/km2 or 8 500 mg/km2.  \n\nRivers: Mean 7.0 (± 0.20) 10-3 #/L.  \n\n"
        "Median: 0.36 10-3 #/L",
        "n.i.",
        "Manta trawl; 300{mu}m mesh",
        "Visual sorting; FTIR on a subset.",
        "",
    ],
    [
        "Simon et al. 2018",
        "DK",
        "WWTP",
        "3° treatment",
        "2017",
        "Size: Up to 600 µm",
        "PE, PP, PEST",
        "Raw wastewater median:: 7216 #/L or 250 ug/L \nTreated wastewater median: 54 #/L or "
        "4.2 ug/L \nRecovery efficiency=99.3%",
        "Blanks and spiked recovery included.",
        " \nRaw WW samples filtered on-site through 10 µm stainless steel meshes.",
        "Enzymatic and oxidative digestion; FPA-µFTIR imaging.",
        "",
    ],
    [
        "Browne et al. 2011",
        "West Hornsby and Hornsby Heights, NSW, AU",
        "WWTP",
        "3° treatment",
        "2010",
        "n.i.",
        "PEST, PMMA and PA",
        "EF mean: 1 #/L",
        "n.i.",
        "Samples  collected in glass bottles with metal caps. ",
        "Filtered and identified with Transmittance FT-IR",
        "\xa0",
    ],
]


def _fill(cell, text):
    paragraph = cell.paragraphs[0]
    for i, part in enumerate(text.split("{mu}")):
        if i:
            sym = OxmlElement("w:sym")
            sym.set(qn("w:font"), "Symbol")
            sym.set(qn("w:char"), "F06D")
            paragraph.add_run()._r.append(sym)
        if part:
            paragraph.add_run(part)


def build_fixture_docx(path):
    """How tests/fixtures/koelmans2019_table_s1_fixture.docx was generated."""
    doc = Document()
    doc.add_paragraph("Table S1: Study Characteristics")
    legend = doc.add_table(rows=2, cols=2)
    legend.rows[0].cells[0].text, legend.rows[0].cells[1].text = "Abbreviation", "Full name"
    legend.rows[1].cells[0].text, legend.rows[1].cells[1].text = "LAK", "Lake"
    table = doc.add_table(rows=1, cols=len(REAL_HEADERS))
    for cell, header in zip(table.rows[0].cells, REAL_HEADERS, strict=True):
        cell.text = header
    for row in FIXTURE_ROWS:
        for cell, text in zip(table.add_row().cells, row, strict=True):
            _fill(cell, text)
    doc.save(path)
    return path


@pytest.fixture
def fixture_docx(fixtures_dir):
    path = fixtures_dir / "koelmans2019_table_s1_fixture.docx"
    assert path.is_file(), "committed fixture missing; regenerate with build_fixture_docx"
    return path


def test_committed_fixture_parses_the_same_as_a_freshly_built_one(tmp_path, fixture_docx):
    fresh = build_fixture_docx(tmp_path / "fresh.docx")
    pd.testing.assert_frame_equal(k.parse_table_s1(fixture_docx), k.parse_table_s1(fresh))


def test_symbol_font_micro_sign_is_recovered(fixture_docx):
    tables = k.read_tables(fixture_docx)
    table = next(t for t in tables if t[0][0] == "Reference")
    sampling = table[1][REAL_HEADERS.index("Sampling method")]
    assert "333 µm mesh" in sampling
    value = table[3][REAL_HEADERS.index("Value")]
    assert "> 100 µm - 10.4 #/L" in value
    assert "\n" in table[1][REAL_HEADERS.index("Value")]


@pytest.mark.parametrize(
    ("raw", "norm"),
    [
        ("Size, shape", "size shape"),
        ("  Reference\n", "reference"),
        ("Size range (µm)", "size range"),
        ("Sampling Method", "sampling method"),
        (
            "Detection limit, negative and positive controls and blanks.  ",
            "detection limit negative and positive controls and blanks",
        ),
    ],
)
def test_normalise_header(raw, norm):
    assert k.normalise_header(raw) == norm


def test_map_headers_on_the_real_header_row():
    roles = k.map_headers(REAL_HEADERS)
    assert roles == {
        "study": 0,
        "source": 2,
        "size": 5,
        "value": 7,
        "sampling": 9,
        "method": 10,
        "notes": 11,
    }


def test_map_headers_is_name_driven_not_position_driven():
    shuffled = list(reversed(REAL_HEADERS))
    roles = k.map_headers(shuffled)
    assert shuffled[roles["study"]] == "Reference"
    assert shuffled[roles["value"]] == "Value"
    assert shuffled[roles["size"]] == "Size, shape"


def test_map_headers_accepts_tidy_synonyms():
    roles = k.map_headers(["Study", "Matrix", "Size range (µm)", "Concentration", "Method"])
    assert roles == {"study": 0, "source": 1, "size": 2, "value": 3, "method": 4}


def test_map_headers_raises_listing_headers_and_missing_roles():
    with pytest.raises(k.HeaderMappingError) as exc:
        k.map_headers(["Foo", "Bar", "Value"])
    msg = str(exc.value)
    assert "Foo" in msg and "Bar" in msg
    assert "study" in msg and "size" in msg


@pytest.mark.parametrize(
    ("text", "expected"),
    [
        ("Size: 0.355-0.999, 1.00-4.759, ≥4.75 mm; \n\nShape: Fragments", (355.0, None)),
        ("Size: >4750, 4750-1000 and 106-1000 µm.", (106.0, None)),
        ("Size: 112-300 µm, 300-500 µm, 500 µm - 1.6 mm, and 1.6-5 mm;", (112.0, 5000.0)),
        ("Size: 50 – 5000 µm.", (50.0, 5000.0)),
        ("Size: 0.112−0.5 mm, 0.5−1 mm, and 1−5 mm;", (112.0, 5000.0)),
        ("Size: 6.5-100, >100 µm;", (6.5, None)),
        ("Size: 5-10µm, 10-20µm, 20-50µm, 50 - 100µm, > 100µm", (5.0, None)),
        ("Size: 1-5, 5-10, 10-50, 50-100, >100 µm;", (1.0, None)),
        ("Size: (20), 45, 180, 400 µm.", (20.0, None)),
        ("Size: 200, 100 and 20µm;", (20.0, None)),
        ("Size: 300-5000 µm and <300 µm.", (None, 5000.0)),
        ("Size: <0.5, 0.5-1, 1-2, 2-3, 3-4 mm;", (None, 4000.0)),
        ("Size: Up to 600 µm", (None, 600.0)),
        ("Fibres (bulk)50-4850 µm. Fragments (net): > 300 µm.", (50.0, 4850.0)),
        ("Size: 0.598 ± 0.089 mm.", (None, None)),
        ("Size IF, median: 50µm  \nSize EF, median: 51.5µm.", (None, None)),
        ("Shape: Foam, film, fibre, fragment, pellet", (None, None)),
        ("n.i.", (None, None)),
    ],
)
def test_parse_size(text, expected):
    assert k.parse_size(text) == expected


@pytest.mark.parametrize(
    ("text", "expected"),
    [
        ("Manta trawl; 333 µm mesh. Preserved in 70% ethanol.", 333.0),
        ("Neuston net (333 µm). Sample volume measured.", 333.0),
        ("Pumped 3 - 200 L through stacked sieves of 500, 190, 100 and 25 µm at 10 L/min.", 25.0),
        ("Raw WW samples filtered on-site through 10 µm stainless steel meshes.", 10.0),
        ("Trawl net, 0.112 mm mesh, volume from flow & net size", 112.0),
        ("Samples collected in glass bottles with metal caps.", None),
        ("Pumped 25L of water from 1m depth (2 reps) and filtered through 48 µm sieve", 48.0),
    ],
)
def test_mesh_floor(text, expected):
    assert k.mesh_floor(text) == expected


@pytest.mark.parametrize(
    ("text", "value", "unit", "kind"),
    [
        (
            "Mean: 193 420 (±115 567 SD) #/km2\nRange: 52 508 – 748 027 #/km2",
            193420,
            "#/km2",
            "mean",
        ),
        (
            "Mean: 4.2 x 10-3 #/L;\nMedian: 1.9 x 10-3 #/”L;\nRange: 0.05 - 32 (x10-3) #/L",
            0.0042,
            "#/L",
            "mean",
        ),
        ("EF mean: 1 #/L", 1.0, "#/L", "mean"),
        ("Mean (s.d.): 4.703 (± 2.816) # /L; Range: 1.597-12.611 #/L", 4.703, "#/L", "mean"),
        (
            "IF: mean 293 (range: 260-320) #/L  \n\nEF: mean 35(range: 14-50) #/L",
            293.0,
            "#/L",
            "mean",
        ),
        ("Plankton net: mean 30 (range 3 – 106) 10-3 #/L.", 0.03, "#/L", "mean"),
        ("Mean  892,777 #/km2; \nRange: 52 364 -3 931 062  #/km2", 892777.0, "#/km2", "mean"),
        ("Mean: 0.7 x 10-3  #/L; Range 0 – 7 × 10-3 #/L", 0.0007, "#/L", "mean"),
        ("Single-use plastic bottles. Mean: 14 ±14 #/L; \nRange: 2 - 44 #/L.", 14.0, "#/L", "mean"),
        (
            "Lakes: Mean 91 000 (±120 000) #/km-2 or 26 000 (33 000) mg/km2.",
            91000.0,
            "#/km-2",
            "mean",
        ),
        ("Raw wastewater median:: 7216 #/L or 250 ug/L", 7216.0, "#/L", "median"),
        ("Median: 0.98 #/L, minimum: 0.625 #/L", 0.98, "#/L", "median"),
        ("Median EF: 0.07 #/L", 0.07, "#/L", "median"),
        ("IF=133.0 ± 35.6 #/L", 133.0, "#/L", "single"),
        ("Average: 0.3 ± 0.1 (S.E.) #/L", 0.3, "#/L", "mean"),
        ("Range: 3.36 - 6.42 (x10-3) #/L", np.sqrt(3.36e-3 * 6.42e-3), "#/L", "range_geomean"),
        (
            "Range: 0.005-0.0517 mg/L and 0.058-1.265 #/L",
            np.sqrt(0.058 * 1.265),
            "#/L",
            "range_geomean",
        ),
        ("0.55 E5 - 342 E5 #/km2", np.sqrt(0.55e5 * 342e5), "#/km2", "range_geomean"),
        (
            "Range lake: 0.05-7.58 E5 particles/km2.",
            np.sqrt(0.05e5 * 7.58e5),
            "particles/km2",
            "range_geomean",
        ),
        (
            "Range: 1.660 ± 0.6391 – 8.925 ± 1.591 #/L.",
            np.sqrt(1.66 * 8.925),
            "#/L",
            "range_geomean",
        ),
        (
            "IF:  \nMedian: 5.9\nMean: 8.0mg/L.  \n\nEF:  \nMedian: 0.016, Mean:  0.034mg/L.  \n\n"
            "IF:  \nMedian: 86000, Mean: 127000#/L.  \nEF:  \nMedian: 6400, Mean: 5800 #/L.  \n",
            127000.0,
            "#/L",
            "mean",
        ),
        ("Mean: 5 mg/L", 5.0, "mg/L", "mean"),
        (
            "Mean for 125-2000 µm: upstream WWTP: 24 (±11.4) 10-3 #/L; downstream WWTP: "
            "71.7 (± 60.2) 10-3 #/L.",
            0.024,
            "#/L",
            "single",
        ),
    ],
)
def test_parse_value(text, value, unit, kind):
    got_value, got_unit, got_kind = k.parse_value(text)
    assert got_value == pytest.approx(value, rel=1e-9)
    assert got_unit == unit
    assert got_kind == kind


@pytest.mark.parametrize(
    "text",
    [
        "n.i.",
        "(1) Skimming Tertiary EF:  3-23 MP in 9.46-9.57 × 106 L skimmed",
        "",
        # the "3" of a folded scale factor is not a concentration
        "Concentration means and ranges at 5 sites (10-3 # / L): 100.6 (5.7-398.0), "
        "48.5 (2.7-441.4).",
    ],
)
def test_parse_value_returns_none_when_no_concentration(text):
    assert k.parse_value(text) == (None, None, "none")


@pytest.mark.parametrize(
    ("code", "matrix"),
    [
        ("LAK", "freshwater"),
        ("RIV", "freshwater"),
        ("LAK, RIV", "freshwater"),
        ("Canal", "freshwater"),
        ("BOT", "drinking water"),
        ("TAP", "drinking water"),
        ("DWTP \n", "drinking water"),
        ("GROUND \n", "groundwater"),
        ("WWTP", "wastewater"),
        ("???", "other"),
    ],
)
def test_matrix_from_source(code, matrix):
    assert k.matrix_from_source(code) == matrix


def test_parse_table_s1_maps_the_fixture_to_the_canonical_columns(fixture_docx):
    df = k.parse_table_s1(fixture_docx)
    assert list(df.columns[: len(k.CANONICAL)]) == list(k.CANONICAL)
    assert len(df) == len(FIXTURE_ROWS)
    by = df.set_index("study")

    anderson = by.loc["Anderson et al. 2017"]
    assert anderson["matrix"] == "freshwater"
    assert anderson["size_min_um"] == 333.0 and np.isnan(anderson["size_max_um"])
    assert anderson["concentration_value"] == 193420.0
    assert anderson["concentration_unit"] == "#/km2"
    assert anderson["method"].startswith("Samples rinsed")
    assert "size_source=sampling_mesh" in anderson["notes"]

    baldwin = by.loc["Baldwin et al., 2016"]
    assert baldwin["size_min_um"] == 355.0 and np.isnan(baldwin["size_max_um"])
    assert baldwin["concentration_value"] == pytest.approx(0.0042)
    assert "size_source=size_text" in baldwin["notes"]

    mason = by.loc["Mason et al. 2018"]
    assert mason["matrix"] == "drinking water"
    assert mason["size_min_um"] == 6.5
    assert mason["concentration_value"] == 325.0

    faure = by.loc["Faure et al. 2015"]
    assert (faure["size_min_um"], faure["size_max_um"]) == (300.0, 5000.0)
    assert faure["concentration_value"] == 91000.0 and faure["concentration_unit"] == "#/km-2"

    simon = by.loc["Simon et al. 2018"]
    assert (simon["size_min_um"], simon["size_max_um"]) == (10.0, 600.0)
    assert simon["concentration_value"] == 7216.0
    assert "value_kind=median" in simon["notes"]

    browne = by.loc["Browne et al. 2011"]
    assert np.isnan(browne["size_min_um"]) and np.isnan(browne["size_max_um"])
    assert browne["concentration_value"] == 1.0


def test_csv_roundtrip_preserves_values_and_missing(tmp_path, fixture_docx):
    df = k.parse_table_s1(fixture_docx)
    path = k.write_table_s1(df, tmp_path / "table_s1.csv")
    back = k.load_table_s1(path)
    pd.testing.assert_frame_equal(back, df)


def _zip_with(docx_path):
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("gr1.gif", b"GIF89a")
        z.writestr("mmc1.docx", docx_path.read_bytes())
    return buf.getvalue()


def test_fetch_si_downloads_once_then_uses_the_cache(tmp_path, fixture_docx):
    payload = _zip_with(fixture_docx)
    calls = []

    def handler(request):
        calls.append(request)
        return httpx.Response(200, content=payload, headers={"content-type": "application/zip"})

    client = httpx.Client(transport=httpx.MockTransport(handler))
    docx_path = k.fetch_si(dest=tmp_path, client=client)
    assert docx_path == tmp_path / "mmc1.docx"
    assert docx_path.read_bytes() == fixture_docx.read_bytes()
    assert (tmp_path / "supplementaryFiles.zip").read_bytes() == payload
    assert len(calls) == 1
    assert str(calls[0].url) == k.SI_URL
    assert calls[0].headers["user-agent"].startswith("NanoSolve-mnp-research/")

    assert k.fetch_si(dest=tmp_path, client=client) == docx_path
    assert len(calls) == 1


def test_fetch_si_fails_loud_on_http_error(tmp_path):
    client = httpx.Client(transport=httpx.MockTransport(lambda r: httpx.Response(404)))
    with pytest.raises(httpx.HTTPStatusError):
        k.fetch_si(dest=tmp_path, client=client)
    assert not (tmp_path / "mmc1.docx").exists()
