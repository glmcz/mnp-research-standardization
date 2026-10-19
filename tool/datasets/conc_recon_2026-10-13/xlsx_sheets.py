import sys

import openpyxl

wb = openpyxl.load_workbook(sys.argv[1], read_only=True, data_only=True)
for ws in wb.worksheets:
    print("==", ws.title, ws.max_row, ws.max_column)
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i >= int(sys.argv[2] if len(sys.argv) > 2 else 4):
            break
        print("  ", [c for c in row[:60]])
